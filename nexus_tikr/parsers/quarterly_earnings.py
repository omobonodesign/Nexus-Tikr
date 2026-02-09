"""Template I: TIKR_QUARTERLY_EARNINGS parser (v2.1 Robust Parsing Protocol)."""

from __future__ import annotations

import logging
import math
import re
from datetime import datetime
from typing import Any, Optional

import openpyxl

from nexus_tikr.constants import QE_SECTION1_LABELS, QE_SECTION2_METRICS
from nexus_tikr.label_engine import SemanticLabelEngine, _parse_bps, _parse_number, _parse_pct_decimal
from nexus_tikr.models import FileInfo, ForwardGroup, LabelHealth, QESectionMap

logger = logging.getLogger(__name__)


def parse_quarterly_earnings(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    file_info: FileInfo,
    label_health: LabelHealth,
    is_latest_revenue: Optional[float] = None,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]], dict[str, Any], list[ForwardGroup], QESectionMap]:
    """Parse Quarterly Earnings file with 3 sections.

    Args:
        ws: Worksheet object.
        file_info: File info.
        label_health: Label health tracker.
        is_latest_revenue: Latest quarterly revenue from IS for cross-validation.

    Returns:
        (section1_data, section2_data, section2_meta, forward_groups, section_map)
    """
    max_row = ws.max_row or 0
    section_map = _detect_section_boundaries(ws, max_row)

    section1_data: dict[str, Any] = {}
    section2_data: dict[str, dict[str, Any]] = {}
    section2_meta: dict[str, Any] = {}
    forward_groups: list[ForwardGroup] = []

    # Parse each section if found
    if section_map.section_1_found:
        section1_data = _parse_section_1(
            ws, section_map.section_1_anchor, section_map.section_1_end,
            label_health,
        )

    if section_map.section_2_found:
        section2_data, section2_meta = _parse_section_2(
            ws, section_map.section_2_anchor, section_map.section_2_end,
            label_health,
        )

    if section_map.section_3_found:
        forward_groups = _parse_section_3(
            ws, section_map.section_3_anchor, section_map.section_3_end,
            label_health, is_latest_revenue,
        )

    return section1_data, section2_data, section2_meta, forward_groups, section_map


def _detect_section_boundaries(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    max_row: int,
) -> QESectionMap:
    """I.0: Detect section boundaries by scanning for anchor patterns."""
    smap = QESectionMap()

    for row in range(1, max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val is None:
            continue
        s = str(val).strip()

        # Section 1: "{TICKER} Q{N} Earnings | TIKR.com"
        if re.match(r"^\w+\s+Q\d+\s+Earnings\s*\|\s*TIKR\.com$", s):
            smap.section_1_anchor = row
            smap.section_1_found = True

        # Section 2: "{TICKER} Beats & Misses | TIKR.com"
        elif re.match(r"^\w+\s+Beats\s*&\s*Misses\s*\|\s*TIKR\.com$", s):
            smap.section_2_anchor = row
            smap.section_2_found = True

        # Section 3: "{TICKER} Analyst Estimates Breakdown | TIKR.com"
        elif re.match(r"^\w+\s+Analyst\s+Estimates\s+Breakdown\s*\|\s*TIKR\.com$", s):
            smap.section_3_anchor = row
            smap.section_3_found = True

    # Calculate section boundaries
    if smap.section_1_found:
        if smap.section_2_found:
            smap.section_1_end = smap.section_2_anchor - 1
        elif smap.section_3_found:
            smap.section_1_end = smap.section_3_anchor - 1
        else:
            smap.section_1_end = max_row

    if smap.section_2_found:
        if smap.section_3_found:
            smap.section_2_end = smap.section_3_anchor - 1
        else:
            smap.section_2_end = max_row

    if smap.section_3_found:
        smap.section_3_end = max_row

    return smap


def _parse_section_1(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    anchor_row: int,
    end_row: int,
    label_health: LabelHealth,
) -> dict[str, Any]:
    """I.1: Parse Section 1 — Latest Quarter Summary."""
    engine = SemanticLabelEngine(ws)
    engine.label_health = label_health

    # Detect column structure from header row (row after anchor)
    header_row = anchor_row + 1
    columns = _detect_section1_columns(ws, header_row)

    result: dict[str, Any] = {}

    # Track EBITDA row for disambiguation
    ebitda_row = None

    for field_name, (label_key, parse_type, impact) in QE_SECTION1_LABELS.items():
        if field_name == "ebitda":
            # Find EBITDA first
            row = engine.find_label(
                label_key,
                file_type="TIKR_QUARTERLY_EARNINGS",
                section_anchor="section_1",
                impact=impact,
                start_row=anchor_row + 2,
                end_row=end_row,
            )
            if row is not None:
                ebitda_row = row
        elif field_name == "ebit":
            # EBIT disambiguation: must NOT be EBITDA row
            row = _find_exact_ebit(engine, anchor_row + 2, end_row, ebitda_row)
        else:
            row = engine.find_label(
                label_key,
                file_type="TIKR_QUARTERLY_EARNINGS",
                section_anchor="section_1",
                impact=impact,
                start_row=anchor_row + 2,
                end_row=end_row,
            )

        if row is not None:
            field_data = {}
            for col_name, col_idx in columns.items():
                raw = ws.cell(row=row, column=col_idx).value
                if parse_type == "mixed_margins":
                    if isinstance(raw, str) and "bps" in raw.lower():
                        field_data[col_name] = _parse_bps(raw)
                    else:
                        field_data[col_name] = _parse_pct_decimal(raw)
                else:
                    field_data[col_name] = _parse_number(raw)
            result[field_name] = field_data

    return result


def _detect_section1_columns(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
) -> dict[str, int]:
    """Detect column positions for Section 1 by reading header text."""
    columns: dict[str, int] = {}

    for col in range(2, (ws.max_column or 10) + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            continue
        s = str(val).strip().lower()

        if "beat" in s or "miss" in s:
            columns["beat_miss"] = col
        elif "yoy" in s:
            columns["yoy_pct"] = col
        elif "qoq" in s:
            columns["qoq_pct"] = col
        elif "street" in s or "estimate" in s or "est" in s:
            columns["estimate"] = col
        elif s in ("a", "e"):
            # A/E marker column — check if it's actual or estimate
            columns.setdefault("actual", col)
        else:
            # Try to detect date columns
            from nexus_tikr.file_access import _parse_date_value
            dt = _parse_date_value(val)
            if dt:
                if "actual" not in columns:
                    columns["actual"] = col
                else:
                    # Could be YoY reference or QoQ reference
                    if "yoy_ref" not in columns:
                        columns["yoy_ref"] = col
                    elif "qoq_ref" not in columns:
                        columns["qoq_ref"] = col

    # If we don't have labeled columns, try positional assignment
    if not columns:
        max_col = min(ws.max_column or 8, 10)
        for col in range(2, max_col + 1):
            val = ws.cell(row=header_row, column=col).value
            if val is not None:
                col_idx = col - 1  # 0-based position
                if col_idx == 1:
                    columns["actual"] = col
                elif col_idx == 2:
                    columns["estimate"] = col
                elif col_idx == 3:
                    columns["yoy_ref"] = col
                elif col_idx == 4:
                    columns["qoq_ref"] = col
                elif col_idx == 5:
                    columns["beat_miss"] = col
                elif col_idx == 6:
                    columns["yoy_pct"] = col
                elif col_idx == 7:
                    columns["qoq_pct"] = col

    return columns


def _find_exact_ebit(
    engine: SemanticLabelEngine,
    start_row: int,
    end_row: int,
    ebitda_row: Optional[int],
) -> Optional[int]:
    """Find EBIT row with exact match, excluding EBITDA rows."""
    for row in range(start_row, end_row + 1):
        raw = engine._col_a_cache.get(row)
        if raw is None:
            continue
        stripped = raw.strip()

        # Exact match for "EBIT" only
        if stripped == "EBIT":
            if ebitda_row is not None and row == ebitda_row:
                continue
            # Verify it doesn't contain EBITDA
            if "EBITDA" not in raw:
                return row

    return None


def _parse_section_2(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    anchor_row: int,
    end_row: int,
    label_health: LabelHealth,
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    """I.2: Parse Section 2 — Beats & Misses History."""
    engine = SemanticLabelEngine(ws)
    engine.label_health = label_health

    # Detect date columns from header
    header_row = anchor_row + 1
    date_columns: dict[str, int] = {}
    for col in range(2, (ws.max_column or 20) + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            continue
        from nexus_tikr.file_access import _parse_date_value
        dt = _parse_date_value(val)
        if dt:
            date_label = dt.strftime("%Y-%m-%d")
            date_columns[date_label] = col

    result: dict[str, dict[str, Any]] = {}

    # Parse metric triplets
    for metric_name, est_label, act_label in QE_SECTION2_METRICS:
        metric_data: dict[str, Any] = {}

        est_row = engine.find_label(
            est_label,
            file_type="TIKR_QUARTERLY_EARNINGS",
            section_anchor="section_2",
            start_row=anchor_row + 2,
            end_row=end_row,
        )
        act_row = engine.find_label(
            act_label,
            file_type="TIKR_QUARTERLY_EARNINGS",
            section_anchor="section_2",
            start_row=anchor_row + 2,
            end_row=end_row,
        )

        if est_row is not None:
            metric_data["estimated"] = engine.extract_row_values(est_row, date_columns, "number")
        if act_row is not None:
            metric_data["actual"] = engine.extract_row_values(act_row, date_columns, "number")

            # Find "% Difference" row immediately after actual
            diff_row = act_row + 1
            diff_val = engine._col_a_cache.get(diff_row, "")
            if "% Difference" in diff_val or "Difference" in diff_val:
                metric_data["surprise_pct"] = engine.extract_row_values(
                    diff_row, date_columns, "pct_decimal"
                )

        if metric_data:
            result[metric_name] = metric_data

    # Find meta rows
    meta: dict[str, Any] = {}

    # Reporting Date
    reporting_row = engine.find_label(
        "Reporting Date",
        file_type="TIKR_QUARTERLY_EARNINGS",
        section_anchor="section_2",
        start_row=anchor_row + 2,
        end_row=end_row,
    )
    if reporting_row:
        meta["reporting_dates"] = engine.extract_row_values(
            reporting_row, date_columns, "raw"
        )

    # 1D Stock Price Change
    price_chg_row = engine.find_label(
        "1D Stock Price Change",
        file_type="TIKR_QUARTERLY_EARNINGS",
        section_anchor="section_2",
        start_row=anchor_row + 2,
        end_row=end_row,
    )
    if price_chg_row:
        meta["price_change_1d"] = engine.extract_row_values(
            price_chg_row, date_columns, "pct_decimal"
        )

    # Calculate track record
    meta["track_record"] = _calc_track_record(result, meta)

    return result, meta


def _calc_track_record(
    data: dict[str, dict[str, Any]],
    meta: dict[str, Any],
) -> dict[str, Any]:
    """Calculate beat/miss track record metrics."""
    record: dict[str, Any] = {}

    for metric_name in ("Revenue", "EBITDA", "Adjusted EPS"):
        metric = data.get(metric_name, {})
        surprises = metric.get("surprise_pct", {})
        if not surprises:
            continue

        vals = [v for v in surprises.values() if v is not None and isinstance(v, (int, float))]
        if not vals:
            continue

        beats = sum(1 for v in vals if v > 0)
        total = len(vals)

        key = metric_name.lower().replace(" ", "_")
        record[f"{key}_beat_rate"] = f"{beats}/{total} ({round(beats / total * 100, 1)}%)"
        record[f"{key}_avg_surprise_pct"] = round(sum(vals) / len(vals), 2)

    # Post-earnings drift
    price_changes = meta.get("price_change_1d", {})
    if price_changes:
        pc_vals = [v for v in price_changes.values()
                   if v is not None and isinstance(v, (int, float))]
        if pc_vals:
            record["avg_post_earnings_drift_pct"] = round(sum(pc_vals) / len(pc_vals), 2)

    # Management credibility
    eps_beat_str = record.get("adjusted_eps_beat_rate", "")
    if eps_beat_str:
        pct_match = re.search(r"\(([\d.]+)%\)", eps_beat_str)
        if pct_match:
            beat_pct = float(pct_match.group(1))
            if beat_pct > 75:
                record["management_credibility"] = "High"
            elif beat_pct >= 50:
                record["management_credibility"] = "Medium"
            else:
                record["management_credibility"] = "Low"

    return record


def _parse_section_3(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    anchor_row: int,
    end_row: int,
    label_health: LabelHealth,
    is_latest_revenue: Optional[float] = None,
) -> list[ForwardGroup]:
    """I.3: Parse Section 3 — Forward Quarterly Estimates with fingerprinting."""
    # Detect columns: Period | Analysts | Average | Median | YoY Chg % | QoQ Chg %
    header_row = anchor_row + 1
    col_map = _detect_section3_columns(ws, header_row)

    if not col_map:
        logger.warning("Could not detect Section 3 column structure")
        return []

    period_col = col_map.get("period", 1)
    analysts_col = col_map.get("analysts")
    avg_col = col_map.get("average")
    median_col = col_map.get("median")
    yoy_col = col_map.get("yoy")
    qoq_col = col_map.get("qoq")

    # Step 1: Extract all rows
    all_rows: list[dict[str, Any]] = []
    for row in range(header_row + 1, end_row + 1):
        period_val = ws.cell(row=row, column=period_col).value
        if period_val is None:
            all_rows.append({"row": row, "period": None, "blank": True})
            continue

        from nexus_tikr.file_access import _parse_date_value
        dt = _parse_date_value(period_val)

        row_data: dict[str, Any] = {
            "row": row,
            "period": dt,
            "period_raw": str(period_val).strip(),
            "blank": False,
        }

        if analysts_col:
            row_data["analysts"] = _parse_number(ws.cell(row=row, column=analysts_col).value)
        if avg_col:
            row_data["average"] = _parse_number(ws.cell(row=row, column=avg_col).value)
        if median_col:
            row_data["median"] = _parse_number(ws.cell(row=row, column=median_col).value)
        if yoy_col:
            row_data["yoy_pct"] = _parse_pct_decimal(ws.cell(row=row, column=yoy_col).value)
        if qoq_col:
            row_data["qoq_pct"] = _parse_pct_decimal(ws.cell(row=row, column=qoq_col).value)

        all_rows.append(row_data)

    # Remove trailing blanks
    while all_rows and all_rows[-1].get("blank"):
        all_rows.pop()

    if not all_rows:
        return []

    # Step 2: Build forward quarter date sequence
    dates = sorted(set(
        r["period"] for r in all_rows
        if r.get("period") is not None and isinstance(r["period"], datetime)
    ))

    if not dates:
        return []

    # Step 3: Detect group boundaries via date reset
    groups: list[list[dict]] = []
    current_group: list[dict] = []
    prev_date = None

    for r in all_rows:
        if r.get("blank"):
            continue
        dt = r.get("period")
        if dt is None:
            continue

        if prev_date is not None and dt <= prev_date:
            # Date reset — new group boundary
            if current_group:
                groups.append(current_group)
            current_group = [r]
        else:
            current_group.append(r)

        prev_date = dt

    if current_group:
        groups.append(current_group)

    # Step 4: Identify metric groups via value-scale fingerprinting
    forward_groups: list[ForwardGroup] = []

    # Calculate average values per group
    group_avgs = []
    for g in groups:
        vals = [r.get("average") for r in g if r.get("average") is not None]
        avg = sum(vals) / len(vals) if vals else 0
        group_avgs.append(abs(avg) if avg != 0 else 0)

    # Expected metric order (by scale, descending): Revenue > EBITDA > EBIT > Adj EPS > EPS GAAP
    metric_names = ["Revenue", "EBITDA", "EBIT", "Adjusted_EPS", "EPS_GAAP"]

    for i, group_rows in enumerate(groups):
        avg_val = group_avgs[i] if i < len(group_avgs) else 0
        magnitude = math.log10(avg_val) if avg_val > 0 else 0

        # Determine metric name
        if i < len(metric_names):
            metric = metric_names[i]
        else:
            metric = f"unidentified_group_{i + 1}"

        # Determine confidence
        confidence = "MEDIUM"
        signals = ["date_reset", f"position_{i + 1}"]

        if is_latest_revenue and i == 0:
            # Cross-validate Revenue
            if abs(avg_val) > 0 and abs(avg_val - is_latest_revenue) / is_latest_revenue < 0.5:
                confidence = "HIGH"
                signals.append("value_scale_match_IS")
            elif magnitude >= 2 and (len(group_avgs) < 2 or avg_val > group_avgs[1] * 1.5 if len(group_avgs) > 1 else True):
                confidence = "HIGH"
                signals.append("largest_values")
        elif i == 0 and len(group_avgs) > 1:
            if avg_val > group_avgs[1] * 1.2:
                confidence = "HIGH"
                signals.append("largest_values")

        # Per-share metrics (magnitude < 2)
        if magnitude < 2 and i >= 3:
            confidence = "MEDIUM"
            signals.append("per_share_scale")

        fg = ForwardGroup(
            metric=metric,
            confidence=confidence,
            signals=signals,
            avg_value=round(avg_val, 2),
            row_start=group_rows[0]["row"] if group_rows else 0,
            row_end=group_rows[-1]["row"] if group_rows else 0,
            rows_data=group_rows,
        )
        forward_groups.append(fg)

    # Validate ordering: Group 1 avg > Group 2 avg > Group 3 avg
    if len(group_avgs) >= 3:
        if not (group_avgs[0] >= group_avgs[1] >= group_avgs[2]):
            for fg in forward_groups[:3]:
                if fg.confidence == "HIGH":
                    fg.confidence = "MEDIUM"

    return forward_groups


def _detect_section3_columns(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    header_row: int,
) -> dict[str, int]:
    """Detect Section 3 column positions from header."""
    col_map: dict[str, int] = {}

    for col in range(1, (ws.max_column or 10) + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            continue
        s = str(val).strip().lower()

        if "period" in s or col == 1:
            col_map.setdefault("period", col)
        elif "analyst" in s or "# of" in s:
            col_map["analysts"] = col
        elif "average" in s or "avg" in s:
            col_map["average"] = col
        elif "median" in s:
            col_map["median"] = col
        elif "yoy" in s:
            col_map["yoy"] = col
        elif "qoq" in s:
            col_map["qoq"] = col

    # Fallback: positional
    if "average" not in col_map:
        cols_with_data = []
        for col in range(2, min((ws.max_column or 8) + 1, 10)):
            val = ws.cell(row=header_row, column=col).value
            if val is not None:
                cols_with_data.append(col)
        if len(cols_with_data) >= 4:
            col_map.setdefault("period", 1)
            col_map.setdefault("analysts", cols_with_data[0])
            col_map.setdefault("average", cols_with_data[1])
            col_map.setdefault("median", cols_with_data[2])
            if len(cols_with_data) >= 5:
                col_map["yoy"] = cols_with_data[3]
                col_map["qoq"] = cols_with_data[4]

    return col_map
