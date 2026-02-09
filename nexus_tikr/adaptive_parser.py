"""Module 9: Adaptive Unknown File Parser."""

from __future__ import annotations

import logging
from typing import Any

import openpyxl

from nexus_tikr.constants import ADAPTIVE_CATEGORIES
from nexus_tikr.label_engine import _parse_number
from nexus_tikr.models import FileInfo

logger = logging.getLogger(__name__)


def parse_adaptive(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    file_info: FileInfo,
) -> dict[str, Any]:
    """Parse an unknown/low-confidence file using adaptive strategy.

    Returns structured extraction result.
    """
    result: dict[str, Any] = {
        "source_file": file_info.filename,
        "confidence": "low",
    }

    # Step 9.1 — Structural Reconnaissance
    recon = _structural_recon(ws)
    result.update(recon)

    # Step 9.2 — Content Classification
    category, cat_confidence = _classify_content(ws)
    result["detected_category"] = category
    result["category_label"] = ADAPTIVE_CATEGORIES[category]["label"]
    result["confidence"] = cat_confidence

    # Step 9.3 — Category-specific extraction
    extraction = _extract_by_category(ws, category, recon)
    result["data"] = extraction
    result["extraction_summary"] = {
        "total_metrics_found": len(extraction) if isinstance(extraction, dict) else 0,
    }

    return result


def _structural_recon(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Step 9.1: Read file structure."""
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    # Check for TIKR origin
    origin = "third_party"
    for row in range(1, min(21, max_row + 1)):
        for col in range(1, min(max_col + 1, 10)):
            val = ws.cell(row=row, column=col).value
            if val and "TIKR.com" in str(val):
                origin = "TIKR"
                break
        if origin == "TIKR":
            break

    # Detect layout type
    # Check if Column A has labels and Row 1 has dates
    col_a_labels = 0
    row_1_data = 0
    for row in range(1, min(21, max_row + 1)):
        val = ws.cell(row=row, column=1).value
        if val and isinstance(val, str) and not str(val).replace(".", "").replace("-", "").isdigit():
            col_a_labels += 1

    for col in range(2, min(max_col + 1, 20)):
        val = ws.cell(row=1, column=col).value
        if val is not None:
            row_1_data += 1

    if col_a_labels > 5 and row_1_data > 3:
        layout = "ROW_ORIENTED"
    elif row_1_data > col_a_labels:
        layout = "COLUMN_ORIENTED"
    else:
        layout = "MIXED"

    # Data density
    total_cells = min(20, max_row) * min(20, max_col)
    non_empty = 0
    for r in range(1, min(21, max_row + 1)):
        for c in range(1, min(21, max_col + 1)):
            if ws.cell(row=r, column=c).value is not None:
                non_empty += 1

    density = non_empty / total_cells if total_cells > 0 else 0

    return {
        "origin_platform": origin,
        "layout_type": layout,
        "data_density": round(density, 2),
        "max_rows": max_row,
        "max_cols": max_col,
    }


def _classify_content(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[str, str]:
    """Step 9.2: Classify file by scanning for financial vocabulary."""
    max_row = ws.max_row or 0
    max_col = min(ws.max_column or 5, 10)

    # Collect all text from the file
    all_text = []
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            val = ws.cell(row=row, column=col).value
            if val and isinstance(val, str):
                all_text.append(val)

    full_text = " ".join(all_text).lower()

    # Score each category
    scores: dict[str, int] = {}
    for cat_key, cat_info in ADAPTIVE_CATEGORIES.items():
        if cat_key == "F":
            continue
        count = sum(1 for trigger in cat_info["triggers"]
                    if trigger.lower() in full_text)
        scores[cat_key] = count

    # Find best match
    best_cat = max(scores, key=scores.get)
    best_score = scores[best_cat]
    threshold = ADAPTIVE_CATEGORIES[best_cat]["threshold"]

    if best_score >= threshold:
        high_threshold = threshold + 2
        if best_score >= high_threshold:
            return best_cat, "high"
        else:
            return best_cat, "medium"

    return "F", "low"


def _extract_by_category(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    category: str,
    recon: dict[str, Any],
) -> dict[str, Any]:
    """Step 9.3: Category-specific extraction."""
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if category in ("A", "B", "C", "D", "E"):
        # Generic label->values extraction
        return _generic_extraction(ws, max_row, max_col)
    else:
        # Category F: raw extraction
        return _raw_extraction(ws, max_row, max_col)


def _generic_extraction(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    max_row: int,
    max_col: int,
) -> dict[str, Any]:
    """Extract all label->values from a row-oriented file."""
    data: dict[str, Any] = {}

    # Try to detect date columns from Row 1
    headers = []
    for col in range(2, max_col + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None:
            headers.append((col, str(val).strip()))

    for row in range(2, max_row + 1):
        label = ws.cell(row=row, column=1).value
        if label is None:
            continue
        label_str = str(label).strip()
        if not label_str:
            continue

        values = {}
        for col, header in headers:
            raw = ws.cell(row=row, column=col).value
            values[header] = _parse_number(raw) if raw is not None else None

        if any(v is not None for v in values.values()):
            data[label_str] = values

    return data


def _raw_extraction(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    max_row: int,
    max_col: int,
) -> dict[str, Any]:
    """Category F: Raw extraction preserving structure."""
    data: dict[str, Any] = {}

    for row in range(1, max_row + 1):
        label = ws.cell(row=row, column=1).value
        if label is None:
            continue
        label_str = str(label).strip()
        if not label_str:
            continue

        values = []
        for col in range(2, max_col + 1):
            val = ws.cell(row=row, column=col).value
            values.append(val)

        data[label_str] = values

    return data
